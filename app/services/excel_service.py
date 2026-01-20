"""
Excel service for saving enriched product data.
"""

import os
import logging
from datetime import datetime
from typing import Dict, Any, Set, List, Optional
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill

from app.core.config import config

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class ExcelService:
    """Service for managing Excel file operations."""
    
    def __init__(self):
        """Initialize Excel service with input path."""
        self.file_path = config.INPUT_PRODUCTS_PATH
        self._ensure_columns_exist()
    
    def _ensure_columns_exist(self) -> None:
        """Ensure enrichment columns exist in the file."""
        if not os.path.exists(self.file_path):
            logger.error(f"Input file not found: {self.file_path}")
            return

        try:
            wb = load_workbook(self.file_path)
            ws = wb.active
            
            # Check headers in first row
            headers = [cell.value for cell in ws[1]]
            required_headers = [
                "Short Description (Enriched)",
                "Product Description (Enriched)", 
                "Product Long Description (Enriched)",
                "Product Division",
                "Class Group",
                "Timestamp"
            ]
            
            # Map of header name to column index (1-based)
            self.header_map = {}
            for idx, header in enumerate(headers, 1):
                if header:
                    self.header_map[header] = idx
            
            # Add missing headers
            new_headers_added = False
            current_max_col = ws.max_column
            header_font = Font(bold=True)

            # Ensure Product Division and Class Group are before Timestamp if they exist
            # Order: ..., Product Division, Class Group, Timestamp
            headers_to_check = ["Product Division", "Class Group"]
            for target_col in headers_to_check:
                if target_col in self.header_map and "Timestamp" in self.header_map:
                    target_idx = self.header_map[target_col]
                    time_idx = self.header_map["Timestamp"]
                    if target_idx > time_idx:
                        logger.info(f"Moving '{target_col}' column before 'Timestamp'")
                        ws.insert_cols(time_idx)
                        for r in range(1, ws.max_row + 1):
                            ws.cell(row=r, column=time_idx).value = ws.cell(row=r, column=target_idx + 1).value
                            if r == 1: ws.cell(row=r, column=time_idx).font = header_font
                        ws.delete_cols(target_idx + 1)
                        # Refresh map
                        headers = [cell.value for cell in ws[1]]
                        self.header_map = {h: i for i, h in enumerate(headers, 1) if h}

            # Special check to ensure Product Division is immediately before Class Group
            if "Product Division" in self.header_map and "Class Group" in self.header_map:
                div_idx = self.header_map["Product Division"]
                class_idx = self.header_map["Class Group"]
                if div_idx != class_idx - 1:
                    logger.info("Ensuring 'Product Division' is next to 'Class Group'")
                    # Move Product Division to class_idx, which shifts Class Group to class_idx + 1
                    ws.insert_cols(class_idx)
                    for r in range(1, ws.max_row + 1):
                        old_col = div_idx if div_idx < class_idx else div_idx + 1
                        ws.cell(row=r, column=class_idx).value = ws.cell(row=r, column=old_col).value
                        if r == 1: ws.cell(row=r, column=class_idx).font = header_font
                    
                    old_col_to_del = div_idx if div_idx < class_idx else div_idx + 1
                    ws.delete_cols(old_col_to_del)
                    
                    # Refresh map
                    headers = [cell.value for cell in ws[1]]
                    self.header_map = {h: i for i, h in enumerate(headers, 1) if h}

            for req in required_headers:
                if req not in self.header_map:
                    # If target is missing but Timestamp exists, insert before it
                    if req in ["Product Division", "Class Group"] and "Timestamp" in self.header_map:
                        col_idx = self.header_map["Timestamp"]
                        ws.insert_cols(col_idx)
                    else:
                        current_max_col += 1
                        col_idx = current_max_col

                    cell = ws.cell(row=1, column=col_idx, value=req)
                    cell.font = header_font
                    self.header_map[req] = col_idx
                    new_headers_added = True
                    
                    # Refresh map to keep indices accurate
                    headers = [cell.value for cell in ws[1]]
                    self.header_map = {h: i for i, h in enumerate(headers, 1) if h}

                    # Set approximate widths
                    if "Description" in req:
                        ws.column_dimensions[cell.column_letter].width = 50
                    elif "Title" in req:
                        ws.column_dimensions[cell.column_letter].width = 30
                    else:
                        ws.column_dimensions[cell.column_letter].width = 20

            if new_headers_added:
                wb.save(self.file_path)
                logger.info("Added missing enrichment columns to Excel file.")
                
        except Exception as e:
            logger.error(f"Failed to ensure columns: {e}")

    def save_enrichment(self, enriched_data: Dict[str, Any]) -> None:
        """
        Save enriched product data to Excel file by updating the row.
        
        Args:
            enriched_data: Dictionary containing enriched product information
        """
        product_code = enriched_data.get("product_code")
        if not product_code:
            raise ValueError("Product code missing in enrichment data")

        try:
            wb = load_workbook(self.file_path)
            ws = wb.active
            
            # Find row for product code
            target_row = None
            code_col_idx = 1 # Assuming Product Code is first column
            
            # Iterate to find the product code
            # Note: For very large files, this might be slow, but fine for typical use
            for row in ws.iter_rows(min_row=2):
                if row[0].value and str(row[0].value).strip() == str(product_code).strip():
                    target_row = row[0].row
                    break
            
            if not target_row:
                raise ValueError(f"Product code {product_code} not found in file")

            # Update validation
            self._ensure_columns_exist() # Re-ensure map is fresh

            # Update cells
            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            
            updates = {
                "Short Description (Enriched)": enriched_data.get("short_title", ""),
                "Product Description (Enriched)": enriched_data.get("short_description", ""),
                "Product Long Description (Enriched)": enriched_data.get("long_description", ""),
                "Product Division": enriched_data.get("product_division", ""),
                "Class Group": enriched_data.get("class_group", ""),
                "Timestamp": timestamp
            }
            
            for header, value in updates.items():
                col_idx = self.header_map.get(header)
                if col_idx:
                    cell = ws.cell(row=target_row, column=col_idx, value=value)
                    if "Description" in header:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")

            wb.save(self.file_path)
            logger.info(f"Updated enrichment for product {product_code}")
            
        except Exception as e:
            logger.error(f"Failed to save to Excel: {e}")
            raise ValueError(f"Excel save failed: {e}")

    def get_existing_product_codes(self) -> Set[str]:
        """
        Get set of product codes that have already been enriched.
        Checks if 'Short Description (Enriched)' column is populated.
        
        Returns:
            Set[str]: Set of existing enriched product codes
        """
        existing_codes = set()
        
        if not os.path.exists(self.file_path):
            return existing_codes
            
        try:
            wb = load_workbook(self.file_path, read_only=True)
            ws = wb.active
            
            # Find "Short Description (Enriched)" column index
            headers = [cell.value for cell in ws[1]]
            short_title_idx = -1
            for idx, header in enumerate(headers):
                if header == "Short Description (Enriched)":
                    short_title_idx = idx
                    break
            
            if short_title_idx == -1:
                return set() # Column doesn't exist, so no products are enriched
            
            # Iterate rows
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and len(row) > short_title_idx:
                    product_code = str(row[0]).strip() if row[0] else ""
                    short_title = str(row[short_title_idx]).strip() if row[short_title_idx] else ""
                    
                    if product_code and short_title:
                         existing_codes.add(product_code)
                    
            logger.info(f"Found {len(existing_codes)} existing enriched products")
            return existing_codes
            
        except Exception as e:
            logger.error(f"Failed to read existing codes: {e}")
            return set()

    def read_products_batch(self, batch_size: int = 10, exclude_codes: Set[str] = None) -> List[Dict[str, str]]:
        """
        Read a batch of products from source file that haven't been enriched yet.
        
        Args:
            batch_size: Number of products to read
            exclude_codes: Set of product codes to skip
            
        Returns:
            List[Dict]: List of products to enrich
        """
        products = []
        exclude_codes = exclude_codes or set()
        
        if not os.path.exists(self.file_path):
            return []
            
        try:
            wb = load_workbook(self.file_path, read_only=True)
            ws = wb.active
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 2:
                    continue
                    
                code = str(row[0]).strip()
                description = str(row[1]).strip() # Assuming description is 2nd column
                
                if not code or not description:
                    continue
                    
                if code in exclude_codes:
                    continue
                    
                products.append({
                    "product_code": code,
                    "product_description": description
                })
                
                if len(products) >= batch_size:
                    break
                    
            logger.info(f"Read {len(products)} new products for enrichment")
            return products
            
        except Exception as e:
            logger.error(f"Failed to read input products: {e}")
            raise ValueError(f"Failed to read input products: {e}")

# Singleton instance
excel_service = ExcelService()
